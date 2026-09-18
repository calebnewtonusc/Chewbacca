#!/usr/bin/env python3
"""Export the entire people db to one wide, readable CSV (and a formatted XLSX)."""
import csv, sqlite3, os, datetime, sys

DB = os.path.expanduser("~/.chewbacca/people/people.db")
OUT_DIR = os.path.expanduser("~/Downloads")
STAMP = datetime.date.today().isoformat()
CSV_PATH = os.path.join(OUT_DIR, f"people-db-{STAMP}.csv")
XLSX_PATH = os.path.join(OUT_DIR, f"people-db-{STAMP}.xlsx")

con = sqlite3.connect(f"file:{DB}?mode=ro", uri=True)
con.row_factory = sqlite3.Row
c = con.cursor()

TODAY = datetime.date.today()

def days_since(ts):
    if not ts:
        return ""
    try:
        d = datetime.datetime.fromisoformat(ts.replace("Z", "")).date()
    except ValueError:
        try:
            d = datetime.datetime.strptime(ts[:10], "%Y-%m-%d").date()
        except ValueError:
            return ""
    return (TODAY - d).days

def clean(v):
    if v is None:
        return ""
    return " ".join(str(v).split())

# ---- aggregates -------------------------------------------------------------
circles = {}
for r in c.execute("""select cm.person_id pid, group_concat(ci.name, ' | ') names
                      from circle_members cm join circles ci on ci.id = cm.circle_id
                      where ci.deleted_at is null group by cm.person_id"""):
    circles[r["pid"]] = r["names"]

msgstats = {}
for r in c.execute("""select person_id pid, count(*) n,
                             sum(from_me) sent,
                             sum(case when from_me=0 then 1 else 0 end) recv,
                             min(sent_at) first, max(sent_at) last
                      from messages where person_id is not null group by person_id"""):
    msgstats[r["pid"]] = r

obscount = {r["pid"]: r["n"] for r in c.execute(
    "select person_id pid, count(*) n from observations where person_id is not null group by person_id")}

intstats = {}
for r in c.execute("""select person_id pid, count(*) n, max(happened_at) last
                      from interactions group by person_id"""):
    intstats[r["pid"]] = r

idents = {}
for r in c.execute("""select person_id pid,
                        group_concat(case when kind='phone' then value end, ' | ') phones,
                        group_concat(case when kind='email' then value end, ' | ') emails
                      from identities group by person_id"""):
    idents[r["pid"]] = r

DIMS = ["spiritual", "emotional", "physical", "intellectual", "social", "financial"]
dimscore = {}
for r in c.execute("""select pds.person_id pid, d.code k, pds.score s
                      from person_dimension_state pds join dimensions d on d.id = pds.dimension_id"""):
    dimscore.setdefault(r["pid"], {})[r["k"]] = r["s"]

# quick facts, with gendered variants folded into one column
FACT_COLS = [
    ("who_they_are", ["who_they_are", "who_he_is", "who_she_is"]),
    ("career",       ["career"]),
    ("works_on",     ["works_on", "works_on_precise"]),
    ("school",       ["school", "college"]),
    ("interests",    ["interests"]),
    ("cares_about",  ["cares_about"]),
    ("faith",        ["faith"]),
    ("family",       ["family"]),
    ("how_they_think", ["how_they_think", "how_he_thinks", "how_she_thinks"]),
    ("advice_they_gave", ["advice_they_gave", "advice_he_gave", "advice_she_gave"]),
    ("ask_about",    ["ask_about"]),
    ("how_to_help",  ["how_to_help"]),
    ("can_introduce", ["can_introduce"]),
    ("first_move",   ["first_move"]),
    ("history",      ["history"]),
    ("music",        ["music"]),
    ("sports",       ["sports"]),
    ("food",         ["food"]),
    ("health",       ["health"]),
    ("_birthday_fact", ["birthday"]),
]
WANTED = {k: col for col, keys in FACT_COLS for k in keys}
facts = {}
for r in c.execute("select person_id, key, value from quick_facts"):
    col = WANTED.get(r["key"])
    if not col:
        continue
    d = facts.setdefault(r["person_id"], {})
    if col in d:
        d[col] += " | " + clean(r["value"])
    else:
        d[col] = clean(r["value"])

# every quick fact key that did NOT map to a column, kept so nothing is lost
other = {}
for r in c.execute("select person_id, key, value from quick_facts"):
    if r["key"] in WANTED:
        continue
    other.setdefault(r["person_id"], []).append(f"{r['key']}: {clean(r['value'])}")

# ---- rows -------------------------------------------------------------------
people = c.execute("""
  select p.*, s.base_score, s.warmth, s.last_interaction_at, s.observation_count, s.completeness
  from people p left join person_scores s on s.person_id = p.id
  where p.deleted_at is null
  order by coalesce(s.base_score,0) desc, p.name collate nocase
""").fetchall()

HEADER = (["rank", "name", "nickname", "score_0_100", "warmth_0_100", "tier", "completeness"]
          + ["company", "role", "location", "linkedin", "phone", "email", "all_phones", "all_emails"]
          + ["circles", "how_we_met", "birthday", "muted", "muted_reason"]
          + ["last_interaction", "days_since_interaction", "interactions_logged",
             "messages_total", "messages_sent", "messages_received", "first_message", "last_message",
             "observations"]
          + [col for col, _ in FACT_COLS if col != "_birthday_fact"]
          + [f"dim_{d}" for d in DIMS]
          + ["other_facts", "source", "created_at", "updated_at", "id"])

def tier(score, days, msgs):
    """Buckets chosen against the real distribution: base_score tops out at 0.43."""
    s = score or 0
    if s >= 0.15: return "core"        # ~105 people
    if s >= 0.08: return "close"       # ~151
    if s >= 0.03: return "known"       # ~221
    if isinstance(days, int) and days > 730: return "dormant"
    if not msgs: return "contact-only"
    return "acquaintance"

rows = []
for i, p in enumerate(people, 1):
    pid = p["id"]
    m = msgstats.get(pid)
    it = intstats.get(pid)
    idn = idents.get(pid)
    f = facts.get(pid, {})
    dims = dimscore.get(pid, {})
    last_i = p["last_interaction_at"] or (m["last"] if m else "") or (it["last"] if it else "")
    ds = days_since(last_i)
    row = [
        i, clean(p["name"]), clean(p["nickname"]),
        round((p["base_score"] or 0) * 100, 1), round((p["warmth"] or 0) * 100, 1),
        tier(p["base_score"], ds, m["n"] if m else 0), round(p["completeness"] or 0, 2),
        clean(p["company"]), clean(p["role"]), clean(p["location"]), clean(p["linkedin"]),
        clean(p["phone"]), clean(p["email"]),
        clean(idn["phones"]) if idn else "", clean(idn["emails"]) if idn else "",
        clean(circles.get(pid)), clean(p["how_we_met"]),
        clean(p["birthday"]) or f.get("_birthday_fact", ""),
        "yes" if p["muted_at"] else "", clean(p["muted_reason"]),
        (last_i or "")[:19], ds, it["n"] if it else 0,
        m["n"] if m else 0, m["sent"] if m else 0, m["recv"] if m else 0,
        (m["first"] or "")[:10] if m else "", (m["last"] or "")[:10] if m else "",
        obscount.get(pid, 0),
    ]
    row += [f.get(col, "") for col, _ in FACT_COLS if col != "_birthday_fact"]
    row += [round(dims.get(d, 0), 1) for d in DIMS]
    row += [" | ".join(other.get(pid, [])), clean(p["source"]),
            (p["created_at"] or "")[:19], (p["updated_at"] or "")[:19], pid]
    rows.append(row)

with open(CSV_PATH, "w", newline="", encoding="utf-8-sig") as fh:
    w = csv.writer(fh)
    w.writerow(HEADER)
    w.writerows(rows)

# ---- xlsx (same data, actually readable when opened) ------------------------
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); ws = wb.active; ws.title = "People"
    ws.append(HEADER); ws.append
    for r in rows: ws.append(r)
    head_fill = PatternFill("solid", fgColor="1F2937")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    widths = {"name": 26, "nickname": 14, "company": 22, "role": 24, "location": 20,
              "linkedin": 30, "phone": 16, "email": 26, "all_phones": 22, "all_emails": 26,
              "circles": 34, "how_we_met": 26, "id": 38}
    for idx, name in enumerate(HEADER, 1):
        w_ = widths.get(name, 40 if name in dict(FACT_COLS) or name == "other_facts" else 13)
        ws.column_dimensions[get_column_letter(idx)].width = w_
    ws.row_dimensions[1].height = 24
    wb.save(XLSX_PATH)
    xlsx_ok = True
except Exception as e:
    xlsx_ok = False
    print("xlsx failed:", e, file=sys.stderr)

print(f"people rows: {len(rows)}  columns: {len(HEADER)}")
print(f"csv : {CSV_PATH}")
if xlsx_ok: print(f"xlsx: {XLSX_PATH}")
